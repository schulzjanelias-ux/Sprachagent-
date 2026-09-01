"""Ersetzt die Mitarbeiternamen der Bauablaufmappe durch erfundene.

Hintergrund: Die Originalmappe der HAG enthaelt 28 echte Mitarbeiternamen.
Als Testgrundlage im Repository wird eine anonymisierte Fassung verwendet
(siehe docs/DECISIONS.md · R-04).

Die Ersetzung erfolgt zellgenau in Listen!B2:B29. Ein globales Suchen und
Ersetzen waere falsch: Der Eintrag 'Berlin' aus der Mitarbeiterliste kommt als
Wort auch im definierten Namen 'Berlin_Feiertage', im Listenkopf und in einem
Zellkommentar vor - ein Textersatz wuerde die Feiertagslogik zerstoeren.

Die Ersatznamen bilden die Eigenheiten des Originals bewusst nach, weil der
Code sie beherrschen muss:
  - ein Nachname doppelt, per Initial unterschieden, mit uneinheitlichem
    Leerzeichen ('S.Dallmann' gegenueber 'D. Dallmann')
  - ein Platzhalter, der keine Person ist ('SO (extern)')
  - Umlaute
  - zwei nachtraeglich angehaengte Eintraege ausserhalb der Sortierung

Aufruf:
    python3 werkzeuge/mappe_anonymisieren.py QUELLE.xlsx ZIEL.xlsx
"""
from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

ERSATZNAMEN = [
    "Ahrens", "Berkhoff", "Böttger", "Cordes",
    "S.Dallmann", "D. Dallmann",          # doppelter Nachname, uneinheitlich geschrieben
    "Emmrich", "Fahnert", "Gollnick", "Grewe", "Hasselbach", "Hüttemann",
    "Immig", "Jarosch", "Kettler", "Lohmeyer", "Mertens", "Osterloh",
    "Petzold", "Reichardt", "Sandbrink", "Terhorst",
    "SO (extern)",                        # Platzhalter, keine Person
    "Uhlig", "Vollmert", "Wesseling",
    "Nowak", "Öztürk",                    # nachtraeglich angehaengt
]

BEREICH_ZEILEN = range(2, 30)   # Listen!B2:B29
BEREICH_SPALTE = 2


def medien_reparieren(original: Path, ziel: Path) -> list[str]:
    """openpyxl verliert eingebettete Medien. Fehlende ZIP-Eintraege zurueckholen."""
    with zipfile.ZipFile(original) as a, zipfile.ZipFile(ziel) as b:
        fehlend = sorted(set(a.namelist()) - set(b.namelist()))
    if not fehlend:
        return []

    zwischen = ziel.with_suffix(".tmp")
    shutil.move(ziel, zwischen)
    with zipfile.ZipFile(original) as quelle, \
            zipfile.ZipFile(zwischen) as alt, \
            zipfile.ZipFile(ziel, "w", zipfile.ZIP_DEFLATED) as neu:
        for eintrag in alt.infolist():
            neu.writestr(eintrag, alt.read(eintrag.filename))
        for name in fehlend:
            neu.writestr(name, quelle.read(name))
    zwischen.unlink()
    return fehlend


def strukturkennzahlen(pfad: Path) -> dict:
    """Kennzahlen, an denen sich eine Beschaedigung der Mappe zeigen wuerde."""
    wb = openpyxl.load_workbook(pfad)
    z = wb["Zeiterfassung"]
    return {
        "blaetter": tuple(wb.sheetnames),
        "merges": len(list(z.merged_cells.ranges)),
        "validierungen": len(list(z.data_validations.dataValidation)),
        "namen": len(list(wb.defined_names)),
        "formel_H6": z["H6"].value,
        "formel_K6": z["K6"].value,
        "diagramme": len(wb["Projektübersicht"]._charts),
        "gewerke": tuple(wb["Listen"].cell(row=r, column=1).value for r in range(2, 14)),
        "feiertage": sum(1 for r in range(2, 83)
                         if wb["Listen"].cell(row=r, column=5).value),
    }


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    quelle, ziel = Path(sys.argv[1]), Path(sys.argv[2])

    if len(ERSATZNAMEN) != len(BEREICH_ZEILEN):
        print(f"Ersatzliste hat {len(ERSATZNAMEN)} Namen, "
              f"erwartet {len(BEREICH_ZEILEN)}", file=sys.stderr)
        return 1

    vorher = strukturkennzahlen(quelle)

    wb = openpyxl.load_workbook(quelle)
    ws = wb["Listen"]
    original_anzahl = sum(1 for r in BEREICH_ZEILEN
                          if ws.cell(row=r, column=BEREICH_SPALTE).value)
    for zeile, ersatz in zip(BEREICH_ZEILEN, ERSATZNAMEN):
        ws.cell(row=zeile, column=BEREICH_SPALTE).value = ersatz
    wb.save(ziel)

    repariert = medien_reparieren(quelle, ziel)

    nachher = strukturkennzahlen(ziel)
    abweichungen = {k: (vorher[k], nachher[k]) for k in vorher if vorher[k] != nachher[k]}

    print(f"Ersetzt : {original_anzahl} Namen in Listen!B2:B29")
    print(f"Medien wiederhergestellt: {repariert or 'keine noetig'}")
    print(f"Struktur unveraendert   : {'ja' if not abweichungen else abweichungen}")
    print(f"  Blaetter {len(nachher['blaetter'])} · Merges {nachher['merges']} · "
          f"Validierungen {nachher['validierungen']} · Namen {nachher['namen']} · "
          f"Gewerke {len(nachher['gewerke'])} · Feiertage {nachher['feiertage']}")

    if abweichungen:
        ziel.unlink()
        print("Ausgabe verworfen.", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
