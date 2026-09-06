"""EPIC 08/09 · Befuellen einer Mappenkopie.

Alle Tests laufen gegen die **echte** Referenzmappe, nicht gegen eine
Nachbildung. Nur so ist belegt, dass Formeln, verbundene Zellen,
Datenvalidierungen und eingebettete Medien den Schreibvorgang ueberstehen
(Briefing §21 I und J, Akzeptanzkriterium 12).
"""
from __future__ import annotations

import datetime
import zipfile
from datetime import date, time
from decimal import Decimal

import openpyxl
import pytest

from app.config import PROJEKTWURZEL
from app.dienste.berichte import Leistung, Tagesbericht, Zeitfenster
from app.dienste.excel.befueller import mappe_befuellen
from app.dienste.excel.leser import MappenFehler, geometrie_pruefen, stammdaten_lesen

REFERENZ = PROJEKTWURZEL / "referenz" / "Bauablaufmappe_Projektcontrolling_anonymisiert.xlsx"
BAUVORHABEN = "Musterstraße 12"
BAUBEGINN = date(2026, 9, 1)
BAUENDE = date(2027, 6, 30)


@pytest.fixture
def projektmappe(tmp_path):
    """Referenzmappe mit gesetzter Bauzeit - so, wie die HAG sie anlegt."""
    ziel = tmp_path / "projektmappe.xlsx"
    mappe = openpyxl.load_workbook(REFERENZ)
    projekt = mappe["Projektübersicht"]
    projekt["B9"] = BAUVORHABEN
    projekt["F11"] = BAUBEGINN
    projekt["F13"] = BAUENDE
    mappe.save(ziel)
    return ziel


def bericht(datum=date(2026, 9, 10), mitarbeiter="Ahrens", gewerk="Trockenbau",
            beginn=time(7, 0), ende=time(16, 30), leistungen=None, kennung="TB-1",
            bemerkung=None) -> Tagesbericht:
    return Tagesbericht(
        bericht_id=kennung, datum=datum, mitarbeiter=mitarbeiter,
        zeitfenster=(Zeitfenster(gewerk, beginn, ende),),
        leistungen=leistungen if leistungen is not None else (
            Leistung("Spachtelarbeiten", menge=Decimal("65"), einheit="m²"),
            Leistung("Schleifarbeiten", menge=Decimal("40"), einheit="m²"),
        ),
        bemerkung=bemerkung,
    )


# --- Glueckspfad -----------------------------------------------------------

def test_bericht_landet_in_der_richtigen_zeile(projektmappe, tmp_path):
    ziel = tmp_path / "befuellt.xlsx"
    plan = mappe_befuellen(projektmappe, ziel, [bericht()],
                           bauvorhaben_erwartet=BAUVORHABEN)
    assert plan.ist_ausfuehrbar and not plan.fehler

    blatt = openpyxl.load_workbook(ziel)["Zeiterfassung"]
    assert blatt["D86"].value == "Ahrens"          # EXCEL-MAPPING §5
    assert blatt["E86"].value == "Trockenbau"
    assert blatt["F86"].value == time(7, 0)
    assert blatt["G86"].value == time(16, 30)
    assert blatt["I86"].value == "Spachtelarbeiten 65 m²; Schleifarbeiten 40 m²"


def test_formelspalten_bleiben_unberuehrt(projektmappe, tmp_path):
    """H darf nie geschrieben werden - Excel rechnet die Stunden (D-03)."""
    ziel = tmp_path / "befuellt.xlsx"
    mappe_befuellen(projektmappe, ziel, [bericht()])
    blatt = openpyxl.load_workbook(ziel)["Zeiterfassung"]
    assert blatt["H86"].value == '=IF(OR(F86="",G86=""),"",ROUND(MOD(G86-F86,1)*24,2))'
    for spalte in ("A", "B", "C", "K"):
        assert str(blatt[f"{spalte}86"].value).startswith("=")


def test_mappe_bleibt_unversehrt(projektmappe, tmp_path):
    """Merges, Validierungen, Diagramm und eingebettete Medien (D-01)."""
    ziel = tmp_path / "befuellt.xlsx"
    mappe_befuellen(projektmappe, ziel, [bericht()])

    vorher = openpyxl.load_workbook(projektmappe)
    nachher = openpyxl.load_workbook(ziel)
    assert nachher.sheetnames == vorher.sheetnames
    assert len(list(nachher["Zeiterfassung"].merged_cells.ranges)) == 1841
    assert len(list(nachher["Zeiterfassung"].data_validations.dataValidation)) == 8
    assert len(list(nachher.defined_names)) == len(list(vorher.defined_names))
    assert len(nachher["Projektübersicht"]._charts) == 1

    with zipfile.ZipFile(projektmappe) as a, zipfile.ZipFile(ziel) as b:
        assert not set(a.namelist()) - set(b.namelist()), "Dateien im Paket verloren"


def test_quellmappe_wird_nicht_veraendert(projektmappe, tmp_path):
    """Der Kern von D-01."""
    vorher = projektmappe.read_bytes()
    mappe_befuellen(projektmappe, tmp_path / "befuellt.xlsx", [bericht()])
    assert projektmappe.read_bytes() == vorher


def test_mehrere_mitarbeiter_und_gewerke_an_einem_tag(projektmappe, tmp_path):
    """Zwei Gewerke eines Mitarbeiters brauchen zwei Zeilen - sonst rechnet
    'Eigenleistung' die Stunden dem falschen Gewerk zu."""
    zwei_gewerke = Tagesbericht(
        bericht_id="TB-2", datum=date(2026, 9, 10), mitarbeiter="Böttger",
        zeitfenster=(Zeitfenster("Trockenbau", time(7, 0), time(12, 0)),
                     Zeitfenster("Maler", time(12, 30), time(16, 0))),
        leistungen=(Leistung("Türen grundieren", menge=Decimal("8"), einheit="Stück"),))

    ziel = tmp_path / "befuellt.xlsx"
    plan = mappe_befuellen(projektmappe, ziel, [bericht(), zwei_gewerke])
    assert len(plan.zuordnungen) == 3

    blatt = openpyxl.load_workbook(ziel)["Zeiterfassung"]
    assert blatt["D87"].value == blatt["D88"].value == "Böttger"
    assert blatt["E87"].value == "Trockenbau"
    assert blatt["E88"].value == "Maler"


def test_leistungen_aller_mitarbeiter_landen_im_selben_tagesfeld(projektmappe, tmp_path):
    """Spalte I ist ueber den Tagesblock verbunden - ein Feld fuer alle."""
    zweiter = bericht(mitarbeiter="Böttger", kennung="TB-2",
                      leistungen=(Leistung("Türen grundieren", menge=Decimal("8"),
                                           einheit="Stück"),))
    ziel = tmp_path / "befuellt.xlsx"
    mappe_befuellen(projektmappe, ziel, [bericht(), zweiter])

    blatt = openpyxl.load_workbook(ziel)["Zeiterfassung"]
    inhalt = blatt["I86"].value
    assert "Spachtelarbeiten 65 m²" in inhalt and "Türen grundieren 8 Stück" in inhalt
    assert blatt["I87"].value is None, "nur die oberste Zelle des Blocks wird beschrieben"


def test_bemerkungen_werden_zusammengefuehrt(projektmappe, tmp_path):
    ziel = tmp_path / "befuellt.xlsx"
    mappe_befuellen(projektmappe, ziel, [
        bericht(bemerkung="Material vollständig"),
        bericht(mitarbeiter="Böttger", kennung="TB-2", bemerkung="Kran defekt")])
    blatt = openpyxl.load_workbook(ziel)["Zeiterfassung"]
    assert blatt["J86"].value == "Material vollständig; Kran defekt"


# --- Vorpruefungen: es darf nichts geschrieben werden ----------------------

def test_probelauf_schreibt_nichts(projektmappe, tmp_path):
    ziel = tmp_path / "befuellt.xlsx"
    plan = mappe_befuellen(projektmappe, ziel, [bericht()], probelauf=True)
    assert plan.zuordnungen and not ziel.exists()


def test_fehlende_bauzeit_bricht_ab(tmp_path):
    """M-1: Ohne Baubeginn erzeugt das Geruest keine Datumswerte."""
    ziel = tmp_path / "befuellt.xlsx"
    plan = mappe_befuellen(REFERENZ, ziel, [bericht()])   # Referenz hat keine Bauzeit
    assert any("Baubeginn" in f for f in plan.fehler)
    assert not ziel.exists()


def test_falsches_bauvorhaben_bricht_ab(projektmappe, tmp_path):
    """D-04: der teuerste denkbare Fehler - Stunden im falschen Projekt."""
    ziel = tmp_path / "befuellt.xlsx"
    plan = mappe_befuellen(projektmappe, ziel, [bericht()],
                           bauvorhaben_erwartet="Hauptstraße 8")
    assert any("Bauvorhaben" in f for f in plan.fehler)
    assert not ziel.exists()


def test_sonntag_wird_abgewiesen(projektmappe, tmp_path):
    plan = mappe_befuellen(projektmappe, tmp_path / "x.xlsx",
                           [bericht(datum=date(2026, 9, 6))])   # Sonntag
    assert any("Sonntag" in f for f in plan.fehler)


def test_feiertag_wird_abgewiesen(projektmappe, tmp_path):
    """D-06: Am 03.10. sperrt die Datenvalidierung der Mappe die Eingabe."""
    plan = mappe_befuellen(projektmappe, tmp_path / "x.xlsx",
                           [bericht(datum=date(2026, 10, 3))])
    assert any("Feiertag" in f for f in plan.fehler)


def test_unbekannter_mitarbeiter_wird_abgewiesen(projektmappe, tmp_path):
    """D-05: Ein abweichender Name faellt lautlos aus der SUMIF-Summe."""
    plan = mappe_befuellen(projektmappe, tmp_path / "x.xlsx",
                           [bericht(mitarbeiter="Mustermann")])
    assert any("Mitarbeiterliste" in f for f in plan.fehler)


def test_unbekanntes_gewerk_wird_abgewiesen(projektmappe, tmp_path):
    plan = mappe_befuellen(projektmappe, tmp_path / "x.xlsx",
                           [bericht(gewerk="Dachdecker")])
    assert any("Gewerkeliste" in f for f in plan.fehler)


def test_elfte_zeile_am_selben_tag_wird_abgewiesen(projektmappe, tmp_path):
    """R-03: Die Mappe fasst zehn Mitarbeiter je Tag."""
    stammdaten = stammdaten_lesen(projektmappe)
    berichte = [bericht(mitarbeiter=name, kennung=f"TB-{i}")
                for i, name in enumerate(stammdaten.mitarbeiter[:11])]
    plan = mappe_befuellen(projektmappe, tmp_path / "x.xlsx", berichte)
    assert any("fasst 10 je Tag" in f or "R-03" in f for f in plan.fehler)


def test_datum_nach_bauende_wird_abgewiesen(projektmappe, tmp_path):
    plan = mappe_befuellen(projektmappe, tmp_path / "x.xlsx",
                           [bericht(datum=date(2027, 7, 15))])
    assert any("Bauende" in f for f in plan.fehler)


def test_bereits_befuellte_zelle_wird_nicht_still_ueberschrieben(projektmappe, tmp_path):
    """M-4: Von Hand eingetragene Werte duerfen nicht verlorengehen."""
    zwischenstand = tmp_path / "erster_lauf.xlsx"
    mappe_befuellen(projektmappe, zwischenstand, [bericht()])

    plan = mappe_befuellen(zwischenstand, tmp_path / "zweiter_lauf.xlsx",
                           [bericht(mitarbeiter="Cordes", kennung="TB-9")])
    assert any("bereits" in f for f in plan.fehler)

    erzwungen = mappe_befuellen(zwischenstand, tmp_path / "erzwungen.xlsx",
                                [bericht(mitarbeiter="Cordes", kennung="TB-9")],
                                ueberschreiben=True)
    assert erzwungen.ist_ausfuehrbar


# --- Schutz gegen Umbau der Mappe (R-01) ----------------------------------

def test_geometriepruefung_erkennt_zerstoerte_stundenformel(projektmappe):
    mappe = openpyxl.load_workbook(projektmappe)
    mappe["Zeiterfassung"]["H6"] = 8.0
    with pytest.raises(MappenFehler, match="Stundenformel"):
        geometrie_pruefen(mappe)


def test_geometriepruefung_erkennt_verschobene_spalten(projektmappe):
    mappe = openpyxl.load_workbook(projektmappe)
    mappe["Zeiterfassung"]["E5"] = "TÄTIGKEIT"
    with pytest.raises(MappenFehler, match="Spaltenueberschriften"):
        geometrie_pruefen(mappe)
