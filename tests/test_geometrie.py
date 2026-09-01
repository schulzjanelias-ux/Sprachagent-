"""EPIC 09 · Adressierung im Blatt 'Zeiterfassung'.

Verifikation auf drei unabhaengigen Wegen:

1. gegen die Formeln der echten Referenzmappe (Spalte K),
2. gegen eine getrennt implementierte Referenz von WORKDAY.INTL(…;"0000001"),
   die schrittweise zaehlt statt zu rechnen,
3. ueber die Umkehrbarkeit datum_zu_zeile -> zeile_zu_datum.

Hinweis: Eine Gegenprobe mit einer echten Tabellenkalkulation waere die
schoenste Bestaetigung. LibreOffice ist in der Entwicklungsumgebung vorhanden,
kann dort aber nicht einmal eine triviale Datei laden - deshalb die drei
unabhaengigen Wege oben statt einer Neuberechnung.
"""
from __future__ import annotations

import re
from datetime import date, time, timedelta

import openpyxl
import pytest

from app.config import PROJEKTWURZEL
from app.dienste.excel.geometrie import (
    ERSTE_DATENZEILE, GeometrieFehler, LETZTE_DATENZEILE, SLOTS_JE_TAG,
    TAGESBLOECKE_MAX, ZEILEN_JE_TAG, arbeitsstunden, datum_zu_zeile,
    ist_erfassungstag, tagesindex, uhrzeit_zu_excel, zeile_zu_datum,
)

MAPPE = PROJEKTWURZEL / "referenz" / "Bauablaufmappe_Projektcontrolling_anonymisiert.xlsx"
BAUBEGINN = date(2026, 9, 1)          # ein Dienstag


def workday_intl_referenz(baubeginn: date, n: int) -> date:
    """Unabhaengige Nachbildung von WORKDAY.INTL(baubeginn-1; n; "0000001").

    Zaehlt Tag fuer Tag und ueberspringt Sonntage - bewusst naiv, damit sie
    nichts mit der geschlossenen Formel in geometrie.py gemeinsam hat.
    """
    tag, verbleibend = baubeginn - timedelta(days=1), n
    while verbleibend:
        tag += timedelta(days=1)
        if tag.weekday() != 6:
            verbleibend -= 1
    return tag


# --- 1. Gegen die echte Mappe ---------------------------------------------

def test_blockstruktur_stimmt_mit_den_formeln_der_mappe():
    """Der n-Parameter in Spalte K muss je Zehnerblock um eins wachsen."""
    ws = openpyxl.load_workbook(MAPPE)["Zeiterfassung"]

    for zeile in range(ERSTE_DATENZEILE, ERSTE_DATENZEILE + 40 * ZEILEN_JE_TAG):
        formel = ws.cell(row=zeile, column=11).value
        treffer = re.search(r'WORKDAY\.INTL\([^,]+,(\d+),"0000001"\)', formel)
        assert treffer, f"K{zeile} hat nicht die erwartete Form"
        n_der_mappe = int(treffer.group(1))
        erwarteter_index = (zeile - ERSTE_DATENZEILE) // ZEILEN_JE_TAG + 1
        assert n_der_mappe == erwarteter_index, f"K{zeile}"


def test_letzte_zeile_entspricht_der_mappe():
    ws = openpyxl.load_workbook(MAPPE)["Zeiterfassung"]
    assert LETZTE_DATENZEILE == 3665
    assert ws.cell(row=LETZTE_DATENZEILE, column=11).value is not None
    assert ws.cell(row=LETZTE_DATENZEILE + 1, column=11).value is None


def test_dokumentiertes_beispiel_aus_excel_mapping():
    """docs/EXCEL-MAPPING.md §5: 10.09.2026 -> Tagesindex 9, Blockstart 86."""
    ziel = datum_zu_zeile(BAUBEGINN, date(2026, 9, 10))
    assert (ziel.tagesindex, ziel.blockstart, ziel.zeile) == (9, 86, 86)
    assert (ziel.mitarbeiter, ziel.gewerk) == ("D86", "E86")
    assert (ziel.anfang, ziel.ende) == ("F86", "G86")
    assert (ziel.leistungen, ziel.bemerkung) == ("I86", "J86")


# --- 2. Gegen die unabhaengige Referenz -----------------------------------

def test_tagesindex_deckt_sich_mit_der_referenz_ueber_alle_bloecke():
    for n in range(1, TAGESBLOECKE_MAX + 1):
        tag = workday_intl_referenz(BAUBEGINN, n)
        assert tagesindex(BAUBEGINN, tag) == n, f"{tag} sollte Index {n} haben"


@pytest.mark.parametrize("wochentag_offset", range(7))
def test_tagesindex_unabhaengig_vom_wochentag_des_baubeginns(wochentag_offset):
    """Beginnt der Bau an einem Sonntag, verschiebt die Mappe auf Montag."""
    beginn = date(2026, 9, 1) + timedelta(days=wochentag_offset)
    for n in range(1, 120):
        tag = workday_intl_referenz(beginn, n)
        assert tagesindex(beginn, tag) == n


def test_sonntage_haben_keine_zeile():
    sonntag = date(2026, 9, 6)
    assert sonntag.weekday() == 6
    assert not ist_erfassungstag(sonntag)
    with pytest.raises(GeometrieFehler, match="Sonntag"):
        datum_zu_zeile(BAUBEGINN, sonntag)


def test_feiertage_sind_keine_erfassungstage():
    feiertag = date(2026, 10, 3)          # Tag der Deutschen Einheit
    assert ist_erfassungstag(feiertag) is True          # ohne Liste: nur Sonntagspruefung
    assert ist_erfassungstag(feiertag, {feiertag}) is False


# --- 3. Umkehrbarkeit ------------------------------------------------------

def test_zeile_und_datum_sind_umkehrbar():
    for n in range(1, TAGESBLOECKE_MAX + 1):
        tag = workday_intl_referenz(BAUBEGINN, n)
        ziel = datum_zu_zeile(BAUBEGINN, tag)
        assert zeile_zu_datum(BAUBEGINN, ziel.blockstart) == tag
        assert zeile_zu_datum(BAUBEGINN, ziel.zeile) == tag


# --- Grenzen ---------------------------------------------------------------

def test_alle_zehn_slots_liegen_im_block():
    ziel_erster = datum_zu_zeile(BAUBEGINN, BAUBEGINN, slot=1)
    assert ziel_erster.zeile == ERSTE_DATENZEILE
    ziel_letzter = datum_zu_zeile(BAUBEGINN, BAUBEGINN, slot=SLOTS_JE_TAG)
    assert ziel_letzter.zeile == ERSTE_DATENZEILE + ZEILEN_JE_TAG - 1
    # Leistungstext bleibt immer im Blockanfang, egal welcher Slot
    assert ziel_letzter.leistungen == ziel_erster.leistungen == "I6"


def test_elfter_mitarbeiter_wird_abgewiesen():
    """R-03: Die Mappe fasst hoechstens zehn Mitarbeiter je Tag."""
    with pytest.raises(GeometrieFehler, match="hoechstens"):
        datum_zu_zeile(BAUBEGINN, BAUBEGINN, slot=SLOTS_JE_TAG + 1)


def test_datum_vor_baubeginn_wird_abgewiesen():
    with pytest.raises(GeometrieFehler, match="vor dem Baubeginn"):
        datum_zu_zeile(BAUBEGINN, BAUBEGINN - timedelta(days=1))


def test_ueberlauf_des_kalendergeruests_wird_gemeldet():
    """R-07: 366 Bloecke reichen fuer rund 14 Monate."""
    letzter = workday_intl_referenz(BAUBEGINN, TAGESBLOECKE_MAX)
    assert datum_zu_zeile(BAUBEGINN, letzter).zeile == LETZTE_DATENZEILE - ZEILEN_JE_TAG + 1

    zu_spaet = workday_intl_referenz(BAUBEGINN, TAGESBLOECKE_MAX + 1)
    with pytest.raises(GeometrieFehler, match="verlaengert"):
        datum_zu_zeile(BAUBEGINN, zu_spaet)


# --- Uhrzeiten und Stunden -------------------------------------------------

@pytest.mark.parametrize("zeitpunkt,erwartet", [
    (time(0, 0), 0.0),
    (time(7, 0), 7 / 24),
    (time(16, 30), 0.6875),
    (time(23, 59), 1439 / 1440),
])
def test_uhrzeit_als_excel_bruchteil(zeitpunkt, erwartet):
    wert = uhrzeit_zu_excel(zeitpunkt)
    assert wert == pytest.approx(erwartet)
    assert 0 <= wert < 1, "Datenvalidierung verlangt 0 <= x < 1"


def test_arbeitsstunden_wie_die_excel_formel():
    assert arbeitsstunden(time(7, 0), time(16, 30)) == 9.5
    assert arbeitsstunden(time(6, 30), time(15, 0)) == 8.5
    assert arbeitsstunden(time(22, 0), time(6, 0)) == 8.0     # MOD: Nachtschicht
    assert arbeitsstunden(time(8, 0), time(8, 0)) == 0.0
