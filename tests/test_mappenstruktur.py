"""EPIC 09 · Regressionstest gegen Aenderungen an der Bauablaufmappe (R-01).

Aendert die HAG die Mappe, schlaegt dieser Test fehl - statt dass der Export
still in falsche Zeilen schreibt. Jede Zusicherung entspricht einer Aussage
aus docs/EXCEL-MAPPING.md.
"""
from __future__ import annotations

import re

import openpyxl
import pytest

from app.config import PROJEKTWURZEL

MAPPE = PROJEKTWURZEL / "referenz" / "Bauablaufmappe_Projektcontrolling_anonymisiert.xlsx"


@pytest.fixture(scope="module")
def mappe():
    return openpyxl.load_workbook(MAPPE)


@pytest.fixture(scope="module")
def zeiterfassung(mappe):
    return mappe["Zeiterfassung"]


def test_erwartete_blaetter(mappe):
    assert mappe.sheetnames == [
        "Projektübersicht", "Zeiterfassung", "Zeiterfassung Export",
        "Materialkosten", "Rechnungsstand", "Rechnungsprüfung",
        "Nachunternehmer", "Eigenleistung", "KPI-Basis", "Listen",
    ]


def test_spaltenueberschriften_in_zeile_5(zeiterfassung):
    erwartet = ["KW", "DATUM", "TAG", "MITARBEITER", "GEWERK", "ARBEITSANFANG",
                "ARBEITSENDE", "ARBEITSSTUNDEN", "ARBEITEN / LEISTUNGEN DES TAGES",
                "TAGESBEMERKUNG", "DATUM INTERN"]
    gefunden = [zeiterfassung.cell(row=5, column=s).value for s in range(1, 12)]
    assert gefunden == erwartet


def test_tagesbloecke_sind_zehn_zeilen_hoch(zeiterfassung):
    """366 Tagesbloecke a 10 Zeilen, erster ab Zeile 6 (EXCEL-MAPPING §2)."""
    tagesmerges = [str(b) for b in zeiterfassung.merged_cells.ranges
                   if re.match(r"^[ABCIJ]\d+:[ABCIJ]\d+$", str(b))]
    starts, hoehen = set(), set()
    for bereich in tagesmerges:
        von, bis = re.match(r"[A-Z]+(\d+):[A-Z]+(\d+)", bereich).groups()
        if int(von) >= 6:
            starts.add(int(von))
            hoehen.add(int(bis) - int(von) + 1)

    assert hoehen == {10}, "Blockhoehe ist nicht mehr 10 Zeilen"
    assert min(starts) == 6
    assert len(starts) == 366
    geordnet = sorted(starts)
    assert {b - a for a, b in zip(geordnet, geordnet[1:])} == {10}


def test_formelspalten_sind_unveraendert(zeiterfassung):
    """H und K muessen Formeln bleiben - die App darf sie nie beschreiben."""
    assert zeiterfassung["H6"].value == '=IF(OR(F6="",G6=""),"",ROUND(MOD(G6-F6,1)*24,2))'
    assert "WORKDAY.INTL" in zeiterfassung["K6"].value
    assert '"0000001"' in zeiterfassung["K6"].value, "Sonntagsmaske geaendert"
    for spalte in ("A", "B", "C"):
        assert str(zeiterfassung[f"{spalte}6"].value).startswith("=")


def test_eingabespalten_sind_frei(zeiterfassung):
    """D, E, F, G, I, J duerfen keine Formeln tragen."""
    for spalte in ("D", "E", "F", "G", "I", "J"):
        wert = zeiterfassung[f"{spalte}6"].value
        assert wert is None or not str(wert).startswith("="), spalte


def test_controlling_haengt_an_gewerk_und_stunden(mappe):
    """Der Befund, auf dem das ganze Mapping beruht (D-02, D-03)."""
    formel = mappe["Eigenleistung"]["F6"].value
    assert "SUMIF" in formel
    assert "Zeiterfassung!$E$6:$E$3665" in formel   # Gewerk
    assert "Zeiterfassung!$H$6:$H$3665" in formel   # Arbeitsstunden


def test_stammdatenlisten(mappe):
    listen = mappe["Listen"]
    gewerke = [listen.cell(row=r, column=1).value for r in range(2, 14)]
    assert len(gewerke) == 12 and all(gewerke)
    assert "Trockenbau" in gewerke and "Maler" in gewerke

    mitarbeiter = [listen.cell(row=r, column=2).value for r in range(2, 30)]
    assert len(mitarbeiter) == 28 and all(mitarbeiter)

    feiertage = [listen.cell(row=r, column=5).value for r in range(2, 83)]
    assert len(feiertage) == 81 and all(feiertage)


def test_validierungen_sperren_feiertage(zeiterfassung):
    regeln = {str(dv.sqref): dv for dv in zeiterfassung.data_validations.dataValidation}
    assert "D6:D3665" in regeln and "MitarbeiterListe" in regeln["D6:D3665"].formula1
    assert "E6:E3665" in regeln and "GewerkeListe" in regeln["E6:E3665"].formula1
    for bereich in ("D6:D3665", "E6:E3665", "I6:I3665", "J6:J3665"):
        assert "Berlin_Feiertage" in regeln[bereich].formula1, bereich


def test_definierte_namen(mappe):
    for name, bereich in [
        ("MitarbeiterListe", "'Listen'!$B$2:$B$29"),
        ("GewerkeListe", "'Listen'!$A$2:$A$13"),
        ("Berlin_Feiertage", "'Listen'!$E$2:$E$82"),
    ]:
        assert mappe.defined_names[name].value == bereich
