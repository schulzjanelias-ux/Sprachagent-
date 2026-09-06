"""EPIC 08 · Strukturierter Export (Briefing §14, Testmatrix I und M).

Der massgebliche strukturierte Export ist XLSX/CSV - Spalte I der
Bauablaufmappe ist die lesbare Beigabe, nicht die Datenquelle (D-02).
"""
from __future__ import annotations

import csv
from datetime import date, time
from decimal import Decimal

import openpyxl
import pytest

from app.dienste.berichte import Leistung, Tagesbericht, Zeitfenster
from app.dienste.excel.export import (
    BLATT_ARBEITSZEITEN, BLATT_BERICHTE, BLATT_LEISTUNGEN, schreiben,
)

BAUVORHABEN = "Musterstraße 12"


@pytest.fixture
def berichte():
    return [
        Tagesbericht(
            bericht_id="TB-1", datum=date(2026, 9, 10), mitarbeiter="Ahrens",
            zeitfenster=(Zeitfenster("Trockenbau", time(7, 0), time(16, 30)),),
            leistungen=(
                Leistung("Spachtelarbeiten", "Wände im Wohnzimmer",
                         Decimal("65"), "m²"),
                Leistung("Schleifarbeiten", None, Decimal("40.5"), "m²",
                         geschaetzt=True)),
            bemerkung="Material vollständig",
            transkript="Heute 65 Quadratmeter gespachtelt."),
        Tagesbericht(
            bericht_id="TB-2", datum=date(2026, 9, 10), mitarbeiter="Böttger",
            zeitfenster=(Zeitfenster("Trockenbau", time(7, 0), time(12, 0)),
                         Zeitfenster("Maler", time(12, 30), time(16, 0))),
            leistungen=(Leistung("Türen grundieren", None, Decimal("8"), "Stück"),)),
    ]


def csv_lesen(pfad):
    with pfad.open(encoding="utf-8-sig", newline="") as datei:
        return list(csv.reader(datei, delimiter=";"))


# --- Aufbau ----------------------------------------------------------------

def test_drei_blaetter_mit_den_erwarteten_zeilen(berichte, tmp_path):
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    mappe = openpyxl.load_workbook(dateien.xlsx)

    assert mappe.sheetnames == [BLATT_BERICHTE, BLATT_LEISTUNGEN, BLATT_ARBEITSZEITEN]
    assert mappe[BLATT_BERICHTE].max_row == 3          # Kopf + 2 Berichte
    assert mappe[BLATT_LEISTUNGEN].max_row == 4        # Kopf + 3 Positionen
    assert mappe[BLATT_ARBEITSZEITEN].max_row == 4     # Kopf + 3 Zeitfenster


def test_getrennte_blaetter_verhindern_doppelte_stunden(berichte, tmp_path):
    """Eine flache Tabelle mit Leistungen UND Zeiten wuerde die Stunden
    mehrfach fuehren - wer die Spalte summiert, bekaeme ein falsches
    Ergebnis. Deshalb getrennt, verbunden ueber die Bericht-Kennung."""
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    blatt = openpyxl.load_workbook(dateien.xlsx)[BLATT_ARBEITSZEITEN]
    stunden = [z[7] for z in blatt.iter_rows(min_row=2, values_only=True)]
    assert sum(stunden) == pytest.approx(9.5 + 5.0 + 3.5)

    berichte_blatt = openpyxl.load_workbook(dateien.xlsx)[BLATT_BERICHTE]
    gesamt = [z[6] for z in berichte_blatt.iter_rows(min_row=2, values_only=True)]
    assert [float(w) for w in gesamt] == [9.5, 8.5]


def test_mengen_bleiben_zahlen_nicht_text(berichte, tmp_path):
    """Damit im Buero gerechnet werden kann, ohne vorher umzuwandeln."""
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    blatt = openpyxl.load_workbook(dateien.xlsx)[BLATT_LEISTUNGEN]
    mengen = [z[7] for z in blatt.iter_rows(min_row=2, values_only=True)]
    assert all(isinstance(w, (int, float, Decimal)) for w in mengen)
    assert [float(w) for w in mengen] == [65.0, 40.5, 8.0]


def test_kopfzeile_ist_fixiert_und_filterbar(berichte, tmp_path):
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    blatt = openpyxl.load_workbook(dateien.xlsx)[BLATT_LEISTUNGEN]
    assert blatt.freeze_panes == "A2"
    assert blatt.auto_filter.ref is not None


# --- CSV -------------------------------------------------------------------

def test_csv_ist_fuer_deutsches_excel_lesbar(berichte, tmp_path):
    """UTF-8 mit BOM und Semikolon: oeffnet ohne Importdialog."""
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    rohbytes = dateien.leistungen_csv.read_bytes()
    assert rohbytes.startswith(b"\xef\xbb\xbf"), "BOM fehlt"
    assert b";" in rohbytes


def test_csv_nutzt_dezimalkomma(berichte, tmp_path):
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    zeilen = csv_lesen(dateien.leistungen_csv)
    mengen = [z[7] for z in zeilen[1:]]
    assert mengen == ["65", "40,5", "8"]


def test_csv_datum_und_uhrzeit_deutsch(berichte, tmp_path):
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    leistungen = csv_lesen(dateien.leistungen_csv)
    assert leistungen[1][1] == "10.09.2026"
    zeiten = csv_lesen(dateien.arbeitszeiten_csv)
    assert zeiten[1][5] == "07:00" and zeiten[1][6] == "16:30"


def test_sonderzeichen_ueberleben(berichte, tmp_path):
    """Briefing §21 M: Umlaute, m², Bauvorhaben mit ß."""
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    text = dateien.leistungen_csv.read_text(encoding="utf-8-sig")
    for zeichen in ("m²", "Böttger", "Musterstraße 12", "Wände im Wohnzimmer",
                    "Türen grundieren", "Stück"):
        assert zeichen in text, zeichen


def test_semikolon_im_freitext_zerstoert_die_csv_nicht(tmp_path):
    heikel = [Tagesbericht(
        bericht_id="TB-9", datum=date(2026, 9, 10), mitarbeiter="Ahrens",
        zeitfenster=(Zeitfenster("Maler", time(7, 0), time(15, 0)),),
        leistungen=(Leistung("Streichen", "Erst grundiert; dann lackiert",
                             Decimal("12"), "m²"),))]
    dateien = schreiben(heikel, tmp_path, bauvorhaben=BAUVORHABEN)
    zeilen = csv_lesen(dateien.leistungen_csv)
    assert len(zeilen) == 2
    assert zeilen[1][6] == "Erst grundiert; dann lackiert"


# --- Reproduzierbarkeit (Briefing §14) ------------------------------------

def test_export_ist_reproduzierbar(berichte, tmp_path):
    erste = schreiben(berichte, tmp_path / "a", bauvorhaben=BAUVORHABEN)
    zweite = schreiben(berichte, tmp_path / "b", bauvorhaben=BAUVORHABEN)
    assert erste.leistungen_csv.read_bytes() == zweite.leistungen_csv.read_bytes()
    assert erste.arbeitszeiten_csv.read_bytes() == zweite.arbeitszeiten_csv.read_bytes()


def test_dateiname_traegt_projekt_und_zeitraum(berichte, tmp_path):
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN,
                        von=date(2026, 9, 1), bis=date(2026, 9, 30))
    assert dateien.xlsx.name == \
        "tagesberichte_Musterstraße_12_2026-09-01_bis_2026-09-30.xlsx"


def test_ohne_zeitraum_bleibt_der_name_eindeutig(berichte, tmp_path):
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    assert "anfang_bis_ende" in dateien.xlsx.name


# --- Grenzfaelle -----------------------------------------------------------

def test_leerer_export_erzeugt_dateien_mit_kopfzeile(tmp_path):
    dateien = schreiben([], tmp_path, bauvorhaben=BAUVORHABEN)
    assert all(p.exists() for p in dateien.alle())
    assert openpyxl.load_workbook(dateien.xlsx)[BLATT_LEISTUNGEN].max_row == 1
    assert len(csv_lesen(dateien.leistungen_csv)) == 1


def test_bericht_ohne_leistungen_erscheint_trotzdem(tmp_path):
    """Sonst verschwaende ein Eintrag lautlos aus dem Export."""
    ohne = [Tagesbericht(bericht_id="TB-0", datum=date(2026, 9, 10),
                         mitarbeiter="Ahrens",
                         zeitfenster=(Zeitfenster("Reinigung", time(7, 0), time(9, 0)),))]
    dateien = schreiben(ohne, tmp_path, bauvorhaben=BAUVORHABEN)
    mappe = openpyxl.load_workbook(dateien.xlsx)
    assert mappe[BLATT_BERICHTE].max_row == 2
    assert mappe[BLATT_ARBEITSZEITEN].max_row == 2
    assert mappe[BLATT_LEISTUNGEN].max_row == 1


def test_nachtschicht_wird_richtig_gerechnet(tmp_path):
    nachts = [Tagesbericht(
        bericht_id="TB-N", datum=date(2026, 9, 10), mitarbeiter="Ahrens",
        zeitfenster=(Zeitfenster("Rückbau", time(22, 0), time(6, 0)),))]
    dateien = schreiben(nachts, tmp_path, bauvorhaben=BAUVORHABEN)
    blatt = openpyxl.load_workbook(dateien.xlsx)[BLATT_ARBEITSZEITEN]
    assert float(list(blatt.iter_rows(min_row=2, values_only=True))[0][7]) == 8.0


def test_rohtranskript_wird_mitgefuehrt(berichte, tmp_path):
    """Briefing §6: raw_transcript als Nachweis - das Audio selbst wird
    nach der Transkription geloescht (D-09)."""
    dateien = schreiben(berichte, tmp_path, bauvorhaben=BAUVORHABEN)
    blatt = openpyxl.load_workbook(dateien.xlsx)[BLATT_BERICHTE]
    erste = list(blatt.iter_rows(min_row=2, values_only=True))[0]
    assert erste[8] == "Heute 65 Quadratmeter gespachtelt."
