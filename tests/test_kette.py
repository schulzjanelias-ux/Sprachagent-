"""Ende-zu-Ende: Datenbank -> Excel-Befueller -> Mappenkopie.

Akzeptanzkriterium 15 aus dem Briefing, soweit ohne Oberflaeche pruefbar.
Ausserdem der Wiederherstellungsweg: Geht eine befuellte Mappe verloren,
muss sie aus der Datenbank neu erzeugbar sein (D-14).
"""
from __future__ import annotations

from datetime import date, time
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy.orm import sessionmaker

from app.ablage import (
    anwesenheiten_setzen, bericht_anlegen, berichte_fuer_export, bestaetigen,
    exportlauf_protokollieren, positionen_setzen,
)
from app.config import PROJEKTWURZEL
from app.datenbank import einrichten, sichern
from app.dienste.berichte import Leistung, Zeitfenster
from app.dienste.excel.befueller import mappe_befuellen
from app.modelle import Mitarbeiter, Projekt

REFERENZ = PROJEKTWURZEL / "referenz" / "Bauablaufmappe_Projektcontrolling_anonymisiert.xlsx"
BAUVORHABEN = "Musterstraße 12"


@pytest.fixture
def projektmappe(tmp_path):
    ziel = tmp_path / "projektmappe.xlsx"
    mappe = openpyxl.load_workbook(REFERENZ)
    projekt = mappe["Projektübersicht"]
    projekt["B9"] = BAUVORHABEN
    projekt["F11"] = date(2026, 9, 1)
    projekt["F13"] = date(2027, 6, 30)
    mappe.save(ziel)
    return ziel


@pytest.fixture
def datenbank(tmp_path, projektmappe):
    maschine = einrichten(tmp_path / "kette.db")
    macher = sessionmaker(bind=maschine, expire_on_commit=False, future=True)
    with macher() as offen:
        ahrens = Mitarbeiter(anmeldename="ahrens", anzeigename="M. Ahrens",
                             excel_name="Ahrens", passwort_hash="x")
        boettger = Mitarbeiter(anmeldename="boettger", anzeigename="K. Böttger",
                               excel_name="Böttger", passwort_hash="x")
        projekt = Projekt(name=BAUVORHABEN, mappe_pfad=str(projektmappe),
                          baubeginn=date(2026, 9, 1), bauende=date(2027, 6, 30))
        offen.add_all([ahrens, boettger, projekt])
        offen.commit()
        yield offen, ahrens, boettger, projekt


def test_von_der_datenbank_in_die_mappe(datenbank, projektmappe, tmp_path):
    sitzung, ahrens, boettger, projekt = datenbank

    erster, _ = bericht_anlegen(sitzung, client_uuid="TB-1", mitarbeiter=ahrens,
                                projekt=projekt, datum=date(2026, 9, 10),
                                rohtranskript="65 Quadratmeter gespachtelt …",
                                bemerkung="Material vollständig")
    positionen_setzen(sitzung, erster, [
        Leistung("Spachtelarbeiten", menge=Decimal("65"), einheit="m²"),
        Leistung("Schleifarbeiten", menge=Decimal("40"), einheit="m²")])
    anwesenheiten_setzen(sitzung, erster,
                         [Zeitfenster("Trockenbau", time(7, 0), time(16, 30))])
    bestaetigen(sitzung, erster)

    zweiter, _ = bericht_anlegen(sitzung, client_uuid="TB-2", mitarbeiter=boettger,
                                 projekt=projekt, datum=date(2026, 9, 10))
    positionen_setzen(sitzung, zweiter, [
        Leistung("Türen grundieren", menge=Decimal("8"), einheit="Stück")])
    anwesenheiten_setzen(sitzung, zweiter, [
        Zeitfenster("Trockenbau", time(7, 0), time(12, 0)),
        Zeitfenster("Maler", time(12, 30), time(16, 0))])
    bestaetigen(sitzung, zweiter)
    sitzung.commit()

    berichte = berichte_fuer_export(sitzung, projekt.id)
    assert len(berichte) == 2

    ziel = tmp_path / "befuellt.xlsx"
    plan = mappe_befuellen(projektmappe, ziel, berichte,
                           bauvorhaben_erwartet=BAUVORHABEN)
    assert plan.ist_ausfuehrbar and not plan.fehler
    assert len(plan.zuordnungen) == 3          # Böttger belegt zwei Zeilen

    blatt = openpyxl.load_workbook(ziel)["Zeiterfassung"]
    assert blatt["D86"].value == "Ahrens"
    assert blatt["D87"].value == blatt["D88"].value == "Böttger"
    assert blatt["E88"].value == "Maler"
    assert blatt["F86"].value == time(7, 0) and blatt["G86"].value == time(16, 30)
    # Spalte I fuehrt die Leistungen aller Mitarbeiter des Tages zusammen
    inhalt = blatt["I86"].value
    assert "Spachtelarbeiten 65 m²" in inhalt and "Türen grundieren 8 Stück" in inhalt
    assert blatt["J86"].value == "Material vollständig"
    # H bleibt die Formel - Excel rechnet die Stunden
    assert str(blatt["H86"].value).startswith("=")

    exportlauf_protokollieren(
        sitzung, projekt_id=projekt.id, ziel_datei=str(ziel),
        zuordnungen=[(z.bericht_id, z.zeile, z.slot) for z in plan.zuordnungen])
    sitzung.commit()

    # Ein zweiter Lauf darf die Stunden nicht verdoppeln (M-6)
    assert berichte_fuer_export(sitzung, projekt.id) == []


def test_mappe_ist_aus_der_datenbank_neu_erzeugbar(datenbank, projektmappe, tmp_path):
    """D-14: Die Datenbank ist die Wahrheit, die Mappe die Ausgabe."""
    sitzung, ahrens, _, projekt = datenbank
    bericht, _ = bericht_anlegen(sitzung, client_uuid="TB-1", mitarbeiter=ahrens,
                                 projekt=projekt, datum=date(2026, 9, 10))
    positionen_setzen(sitzung, bericht,
                      [Leistung("Spachtelarbeiten", menge=Decimal("65"), einheit="m²")])
    anwesenheiten_setzen(sitzung, bericht,
                         [Zeitfenster("Trockenbau", time(7, 0), time(16, 30))])
    bestaetigen(sitzung, bericht)
    sitzung.commit()

    berichte = berichte_fuer_export(sitzung, projekt.id, ohne_bereits_exportierte=False)
    erste = mappe_befuellen(projektmappe, tmp_path / "a.xlsx", berichte)
    zweite = mappe_befuellen(projektmappe, tmp_path / "b.xlsx", berichte)

    assert [(z.zeile, z.slot) for z in erste.zuordnungen] == \
           [(z.zeile, z.slot) for z in zweite.zuordnungen]
    blatt_a = openpyxl.load_workbook(tmp_path / "a.xlsx")["Zeiterfassung"]
    blatt_b = openpyxl.load_workbook(tmp_path / "b.xlsx")["Zeiterfassung"]
    for zelle in ("D86", "E86", "F86", "G86", "I86"):
        assert blatt_a[zelle].value == blatt_b[zelle].value


def test_sicherung_und_wiederherstellung(datenbank, tmp_path):
    """VACUUM INTO statt Dateikopie: konsistent auch im laufenden Betrieb."""
    sitzung, ahrens, _, projekt = datenbank
    bericht, _ = bericht_anlegen(sitzung, client_uuid="TB-1", mitarbeiter=ahrens,
                                 projekt=projekt, datum=date(2026, 9, 10))
    positionen_setzen(sitzung, bericht,
                      [Leistung("Spachtelarbeiten", menge=Decimal("12.5"), einheit="m²")])
    sitzung.commit()

    sicherungsdatei = sichern(tmp_path / "sicherung" / "stand.db")
    assert sicherungsdatei.exists()

    from sqlalchemy import create_engine, select
    from app.modelle import Bericht
    zurueck = create_engine(f"sqlite:///{sicherungsdatei}", future=True)
    with sessionmaker(bind=zurueck, future=True)() as offen:
        wieder = offen.scalar(select(Bericht).where(Bericht.client_uuid == "TB-1"))
        assert wieder is not None
        assert wieder.positionen[0].menge == Decimal("12.5")
