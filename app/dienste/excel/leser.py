"""Stammdaten aus der Bauablaufmappe lesen.

Die Mappe ist die Wahrheit fuer Mitarbeiter, Gewerke und Feiertage (T-01).
Doppelte Pflege in der App waere die schlechtere Loesung: Weicht ein Name ab,
faellt er in 'Eigenleistung' lautlos aus der SUMIF-Summe (D-05).

Ausserdem wird hier die Geometrie geprueft, bevor irgendetwas geschrieben
wird - aendert die HAG die Mappe, bricht der Export ab, statt in falsche
Zeilen zu schreiben (R-01).
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl

from app.dienste.excel.geometrie import (
    ERSTE_DATENZEILE, LETZTE_DATENZEILE, TAGESBLOECKE_MAX, ZEILEN_JE_TAG,
)

BLATT_ZEITERFASSUNG = "Zeiterfassung"
BLATT_PROJEKT = "Projektübersicht"
BLATT_LISTEN = "Listen"

ZELLE_BAUVORHABEN = "B9"
ZELLE_BAUBEGINN = "F11"
ZELLE_BAUENDE = "F13"

SPALTE_GEWERKE, ZEILEN_GEWERKE = 1, range(2, 14)          # Listen!A2:A13
SPALTE_MITARBEITER, ZEILEN_MITARBEITER = 2, range(2, 30)  # Listen!B2:B29
SPALTE_FEIERTAGE, ZEILEN_FEIERTAGE = 5, range(2, 83)      # Listen!E2:E82

FORMEL_STUNDEN = '=IF(OR(F6="",G6=""),"",ROUND(MOD(G6-F6,1)*24,2))'


class MappenFehler(ValueError):
    """Die Mappe ist nicht so aufgebaut, wie der Export es voraussetzt."""


@dataclass(frozen=True)
class Stammdaten:
    bauvorhaben: str | None
    baubeginn: date | None
    bauende: date | None
    gewerke: tuple[str, ...]
    mitarbeiter: tuple[str, ...]
    feiertage: frozenset[date]

    @property
    def bauzeit_gesetzt(self) -> bool:
        return self.baubeginn is not None and self.bauende is not None


def _als_datum(wert) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    return wert if isinstance(wert, date) else None


def _spalte_lesen(blatt, spalte: int, zeilen: range) -> tuple[str, ...]:
    werte = [blatt.cell(row=z, column=spalte).value for z in zeilen]
    return tuple(str(w).strip() for w in werte if w not in (None, ""))


def geometrie_pruefen(mappe) -> None:
    """Gilt die in docs/EXCEL-MAPPING.md dokumentierte Struktur noch?

    Wird vor jedem Schreibvorgang aufgerufen. Lieber ein Abbruch mit Klartext
    als Stunden in fremden Tagen.
    """
    fehlend = [b for b in (BLATT_ZEITERFASSUNG, BLATT_PROJEKT, BLATT_LISTEN)
               if b not in mappe.sheetnames]
    if fehlend:
        raise MappenFehler(f"Blaetter fehlen in der Mappe: {', '.join(fehlend)}")

    blatt = mappe[BLATT_ZEITERFASSUNG]

    kopf = [blatt.cell(row=5, column=s).value for s in range(1, 12)]
    erwartet = ["KW", "DATUM", "TAG", "MITARBEITER", "GEWERK", "ARBEITSANFANG",
                "ARBEITSENDE", "ARBEITSSTUNDEN", "ARBEITEN / LEISTUNGEN DES TAGES",
                "TAGESBEMERKUNG", "DATUM INTERN"]
    if kopf != erwartet:
        raise MappenFehler(
            "Die Spaltenueberschriften in Zeile 5 weichen ab. Erwartet wurde "
            f"{erwartet}, gefunden {kopf}. Die Mappe wurde umgebaut - das "
            "Mapping in docs/EXCEL-MAPPING.md muss nachgezogen werden.")

    if blatt["H6"].value != FORMEL_STUNDEN:
        raise MappenFehler(
            "Die Stundenformel in H6 ist nicht mehr die erwartete. Sie darf "
            "nie ueberschrieben werden - ohne sie bleibt das Controlling leer.")

    formel_k = blatt["K6"].value
    if not isinstance(formel_k, str) or '"0000001"' not in formel_k:
        raise MappenFehler(
            "Das Kalendergeruest in Spalte K nutzt nicht mehr die Sonntagsmaske "
            '"0000001". Die Zuordnung Datum -> Zeile waere damit falsch.')

    # Tagesbloecke: 366 Stueck, je 10 Zeilen, ab Zeile 6
    bloecke = {int(m.group(1))
               for bereich in blatt.merged_cells.ranges
               if (m := re.match(r"^[ABCIJ](\d+):[ABCIJ]\d+$", str(bereich)))
               and int(m.group(1)) >= ERSTE_DATENZEILE}
    if len(bloecke) != TAGESBLOECKE_MAX:
        raise MappenFehler(
            f"Erwartet wurden {TAGESBLOECKE_MAX} Tagesbloecke, gefunden "
            f"{len(bloecke)}. Die Blockstruktur hat sich geaendert.")
    geordnet = sorted(bloecke)
    if geordnet[0] != ERSTE_DATENZEILE or \
            {b - a for a, b in zip(geordnet, geordnet[1:])} != {ZEILEN_JE_TAG}:
        raise MappenFehler(
            f"Die Tagesbloecke liegen nicht mehr im Abstand {ZEILEN_JE_TAG} "
            f"ab Zeile {ERSTE_DATENZEILE}.")
    if geordnet[-1] + ZEILEN_JE_TAG - 1 != LETZTE_DATENZEILE:
        raise MappenFehler(
            f"Der Datenbereich endet nicht in Zeile {LETZTE_DATENZEILE}.")


def stammdaten_aus_mappe(mappe, geometrie: bool = True) -> Stammdaten:
    """Stammdaten aus einer bereits geladenen Mappe.

    Getrennt von stammdaten_lesen, damit der Befueller die Mappe nur einmal
    laden muss - bei 3665 Zeilen mit Formeln kostet jedes Laden spuerbar Zeit.
    """
    if geometrie:
        geometrie_pruefen(mappe)

    projekt, listen = mappe[BLATT_PROJEKT], mappe[BLATT_LISTEN]
    bauvorhaben = projekt[ZELLE_BAUVORHABEN].value

    feiertage = frozenset(
        d for z in ZEILEN_FEIERTAGE
        if (d := _als_datum(listen.cell(row=z, column=SPALTE_FEIERTAGE).value))
    )

    stammdaten = Stammdaten(
        bauvorhaben=str(bauvorhaben).strip() if bauvorhaben else None,
        baubeginn=_als_datum(projekt[ZELLE_BAUBEGINN].value),
        bauende=_als_datum(projekt[ZELLE_BAUENDE].value),
        gewerke=_spalte_lesen(listen, SPALTE_GEWERKE, ZEILEN_GEWERKE),
        mitarbeiter=_spalte_lesen(listen, SPALTE_MITARBEITER, ZEILEN_MITARBEITER),
        feiertage=feiertage,
    )

    if not stammdaten.gewerke:
        raise MappenFehler("Die Gewerkeliste (Listen!A2:A13) ist leer.")
    if not stammdaten.mitarbeiter:
        raise MappenFehler("Die Mitarbeiterliste (Listen!B2:B29) ist leer.")
    return stammdaten


def stammdaten_lesen(pfad: Path, geometrie: bool = True) -> Stammdaten:
    """Liest die Mappe read-only aus.

    data_only=False: Die Geometriepruefung braucht die Formeln in H und K.
    Stammdaten (Namen, Datumswerte, Bauvorhaben) sind ohnehin feste Werte und
    kommen so unveraendert durch. Mit data_only=True kaeme aus einer nie
    berechneten Vorlage ueberall None zurueck.
    """
    if not pfad.exists():
        raise MappenFehler(f"Mappe nicht gefunden: {pfad}")

    mappe = openpyxl.load_workbook(pfad, data_only=False)
    try:
        return stammdaten_aus_mappe(mappe, geometrie=geometrie)
    finally:
        mappe.close()
