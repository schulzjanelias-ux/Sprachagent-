"""Strukturierter Export als XLSX und CSV (EPIC 08).

Briefing §14: "Der Export darf keine unstrukturierten Freitexte als alleinige
Datenquelle enthalten." Genau das ist der Zweck dieser Dateien. Spalte I der
Bauablaufmappe traegt die Leistungen als lesbaren Text (D-02, Variante A) -
massgeblich strukturiert sind aber diese Ausgaben hier, mit Menge und Einheit
in eigenen Feldern.

Aufteilung in drei Blaetter statt einer breiten Tabelle: Ein Bericht hat
mehrere Leistungspositionen **und** mehrere Zeitfenster. In einer flachen
Tabelle mit beidem stuenden die Stunden mehrfach - und die erste Person, die
die Spalte aufsummiert, bekaeme ein falsches Ergebnis. Die Bericht-Kennung
verbindet die Blaetter.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.dienste.berichte import Tagesbericht
from app.dienste.excel.geometrie import arbeitsstunden
from app.dienste.leistungstext import menge_deutsch

KOPF_FUELLUNG = PatternFill("solid", fgColor="1F3864")
KOPF_SCHRIFT = Font(color="FFFFFF", bold=True)

BLATT_BERICHTE = "Berichte"
BLATT_LEISTUNGEN = "Leistungen"
BLATT_ARBEITSZEITEN = "Arbeitszeiten"

SPALTEN_BERICHTE = ["Bericht", "Datum", "KW", "Wochentag", "Mitarbeiter",
                    "Bauvorhaben", "Stunden gesamt", "Bemerkung", "Rohtext"]
SPALTEN_LEISTUNGEN = ["Bericht", "Datum", "Mitarbeiter", "Bauvorhaben",
                      "Position", "Tätigkeit", "Beschreibung", "Menge",
                      "Einheit", "Geschätzt"]
SPALTEN_ARBEITSZEITEN = ["Bericht", "Datum", "Mitarbeiter", "Bauvorhaben",
                         "Gewerk", "Arbeitsanfang", "Arbeitsende", "Stunden"]

BREITEN = {"Bericht": 16, "Datum": 12, "Mitarbeiter": 18, "Bauvorhaben": 24,
           "Tätigkeit": 28, "Beschreibung": 32, "Bemerkung": 32, "Rohtext": 60,
           "Gewerk": 18, "Wochentag": 12, "Stunden gesamt": 14}


@dataclass(frozen=True)
class Exportdateien:
    xlsx: Path
    leistungen_csv: Path
    arbeitszeiten_csv: Path

    def alle(self) -> list[Path]:
        return [self.xlsx, self.leistungen_csv, self.arbeitszeiten_csv]


def _dateiname(projekt: str, von: date | None, bis: date | None) -> str:
    """Reproduzierbar und ohne Sonderzeichen (Briefing §14)."""
    sicher = "".join(z if z.isalnum() else "_" for z in projekt).strip("_")
    zeitraum = f"{von or 'anfang'}_bis_{bis or 'ende'}"
    return f"tagesberichte_{sicher}_{zeitraum}"


def _stunden(bericht: Tagesbericht) -> Decimal:
    summe = sum(arbeitsstunden(f.beginn, f.ende) for f in bericht.zeitfenster)
    return Decimal(str(round(summe, 2)))


def _zeilen_berichte(berichte: list[Tagesbericht], bauvorhaben: str) -> list[list]:
    wochentage = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag",
                  "Samstag", "Sonntag"]
    return [[b.bericht_id, b.datum, b.datum.isocalendar().week,
             wochentage[b.datum.weekday()], b.mitarbeiter, bauvorhaben,
             _stunden(b), b.bemerkung or "", b.transkript or ""]
            for b in berichte]


def _zeilen_leistungen(berichte: list[Tagesbericht], bauvorhaben: str) -> list[list]:
    zeilen = []
    for bericht in berichte:
        for nummer, leistung in enumerate(bericht.leistungen, start=1):
            zeilen.append([
                bericht.bericht_id, bericht.datum, bericht.mitarbeiter, bauvorhaben,
                nummer, leistung.taetigkeit, leistung.beschreibung or "",
                leistung.menge, leistung.einheit or "",
                "ja" if leistung.geschaetzt else "nein"])
    return zeilen


def _zeilen_arbeitszeiten(berichte: list[Tagesbericht], bauvorhaben: str) -> list[list]:
    return [[bericht.bericht_id, bericht.datum, bericht.mitarbeiter, bauvorhaben,
             fenster.gewerk, fenster.beginn, fenster.ende,
             Decimal(str(arbeitsstunden(fenster.beginn, fenster.ende)))]
            for bericht in berichte for fenster in bericht.zeitfenster]


def _blatt_fuellen(blatt, spalten: list[str], zeilen: list[list]) -> None:
    blatt.append(spalten)
    for zelle in blatt[1]:
        zelle.fill, zelle.font = KOPF_FUELLUNG, KOPF_SCHRIFT
        zelle.alignment = Alignment(vertical="center")
    blatt.freeze_panes = "A2"
    blatt.auto_filter.ref = f"A1:{get_column_letter(len(spalten))}1"

    for zeile in zeilen:
        blatt.append(zeile)

    for nummer, name in enumerate(spalten, start=1):
        buchstabe = get_column_letter(nummer)
        blatt.column_dimensions[buchstabe].width = BREITEN.get(name, 14)
        if name in ("Datum",):
            for zelle in blatt[buchstabe][1:]:
                zelle.number_format = "dd.mm.yyyy"
        elif name in ("Arbeitsanfang", "Arbeitsende"):
            for zelle in blatt[buchstabe][1:]:
                zelle.number_format = "hh:mm"
        elif name in ("Menge", "Stunden", "Stunden gesamt"):
            for zelle in blatt[buchstabe][1:]:
                zelle.number_format = "0.00"


def _csv_schreiben(pfad: Path, spalten: list[str], zeilen: list[list]) -> None:
    """UTF-8 mit BOM und Semikolon - so oeffnet Excel die Datei im deutschen
    Sprachraum ohne Importdialog. Dezimaltrenner ist das Komma.
    """
    with pfad.open("w", encoding="utf-8-sig", newline="") as datei:
        schreiber = csv.writer(datei, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        schreiber.writerow(spalten)
        for zeile in zeilen:
            schreiber.writerow([_als_text(wert) for wert in zeile])


def _als_text(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, Decimal):
        return menge_deutsch(wert)
    if isinstance(wert, datetime):
        return wert.strftime("%d.%m.%Y %H:%M")
    if isinstance(wert, date):
        return wert.strftime("%d.%m.%Y")
    if hasattr(wert, "strftime"):                 # datetime.time
        return wert.strftime("%H:%M")
    return str(wert)


def schreiben(berichte: list[Tagesbericht], verzeichnis: Path, *,
              bauvorhaben: str, von: date | None = None,
              bis: date | None = None) -> Exportdateien:
    """Erzeugt XLSX und die beiden CSV-Dateien."""
    verzeichnis.mkdir(parents=True, exist_ok=True)
    basis = _dateiname(bauvorhaben, von, bis)

    zeilen = {
        BLATT_BERICHTE: (SPALTEN_BERICHTE, _zeilen_berichte(berichte, bauvorhaben)),
        BLATT_LEISTUNGEN: (SPALTEN_LEISTUNGEN, _zeilen_leistungen(berichte, bauvorhaben)),
        BLATT_ARBEITSZEITEN: (SPALTEN_ARBEITSZEITEN,
                              _zeilen_arbeitszeiten(berichte, bauvorhaben)),
    }

    mappe = openpyxl.Workbook()
    mappe.remove(mappe.active)
    for name, (spalten, daten) in zeilen.items():
        _blatt_fuellen(mappe.create_sheet(name), spalten, daten)

    dateien = Exportdateien(
        xlsx=verzeichnis / f"{basis}.xlsx",
        leistungen_csv=verzeichnis / f"{basis}_leistungen.csv",
        arbeitszeiten_csv=verzeichnis / f"{basis}_arbeitszeiten.csv",
    )
    mappe.save(dateien.xlsx)
    _csv_schreiben(dateien.leistungen_csv, *zeilen[BLATT_LEISTUNGEN])
    _csv_schreiben(dateien.arbeitszeiten_csv, *zeilen[BLATT_ARBEITSZEITEN])
    return dateien
