"""Eine Kopie der Bauablaufmappe zellgenau befuellen.

Grundsatz D-01: Die Meistermappe wird nie beschrieben. Geschrieben wird immer
in eine Kopie, und zwar ausschliesslich in die sechs Eingabespalten:

    D/E/F/G  je Mitarbeiterzeile     I/J  einmal je Tagesblock

Niemals nach A, B, C, H, K oder L - das sind Formeln. Insbesondere H:
Wuerde man Stunden direkt hineinschreiben, waere die Formel weg und das Blatt
fuer alle Folgetage kaputt (D-03).

Ablauf: pruefen -> planen -> schreiben -> Medien reparieren -> nachpruefen.
Schlaegt die Nachpruefung fehl, wird die Ausgabedatei verworfen statt
ausgeliefert.
"""
from __future__ import annotations

import shutil
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

from app.dienste.berichte import Tagesbericht
from app.dienste.excel.geometrie import (
    GeometrieFehler, SLOTS_JE_TAG, Zielzellen, datum_zu_zeile, ist_erfassungstag,
    uhrzeit_zu_excel,
)
from app.dienste.excel.leser import (
    BLATT_PROJEKT, BLATT_ZEITERFASSUNG, MappenFehler, Stammdaten,
    stammdaten_aus_mappe,
)
from app.dienste.leistungstext import serialisieren


@dataclass(frozen=True)
class Schreibvorgang:
    zelle: str
    wert: Any
    zweck: str


@dataclass(frozen=True)
class Zuordnung:
    """Wohin ein Bericht geschrieben wurde - fuer das Exportprotokoll (M-6)."""

    bericht_id: str
    datum: date
    mitarbeiter: str
    zeile: int
    slot: int


@dataclass
class Befuellungsplan:
    schreibvorgaenge: list[Schreibvorgang] = field(default_factory=list)
    zuordnungen: list[Zuordnung] = field(default_factory=list)
    fehler: list[str] = field(default_factory=list)
    warnungen: list[str] = field(default_factory=list)

    @property
    def ist_ausfuehrbar(self) -> bool:
        return not self.fehler and bool(self.schreibvorgaenge)

    @property
    def betroffene_tage(self) -> int:
        return len({z.datum for z in self.zuordnungen})

    def bericht(self) -> str:
        zeilen = []
        if self.fehler:
            zeilen.append(f"FEHLER ({len(self.fehler)}) — es wird nichts geschrieben:")
            zeilen += [f"  · {f}" for f in self.fehler]
        if self.warnungen:
            zeilen.append(f"Warnungen ({len(self.warnungen)}):")
            zeilen += [f"  · {w}" for w in self.warnungen]
        if self.zuordnungen:
            zeilen.append(f"Geplant: {len(self.zuordnungen)} Zeilen "
                          f"an {self.betroffene_tage} Tagen")
            for z in sorted(self.zuordnungen, key=lambda x: (x.datum, x.slot)):
                zeilen.append(f"  {z.datum:%d.%m.%Y}  Zeile {z.zeile:>4}  "
                              f"Slot {z.slot}  {z.mitarbeiter}")
        if not self.zuordnungen and not self.fehler:
            zeilen.append("Nichts zu schreiben.")
        return "\n".join(zeilen)


def _kennzahlen(mappe) -> dict:
    """Kennzahlen, an denen sich eine Beschaedigung der Mappe zeigen wuerde."""
    blatt = mappe[BLATT_ZEITERFASSUNG]
    return {
        "blaetter": tuple(mappe.sheetnames),
        "merges": len(list(blatt.merged_cells.ranges)),
        "validierungen": len(list(blatt.data_validations.dataValidation)),
        "namen": len(list(mappe.defined_names)),
        "formel_H6": blatt["H6"].value,
        "formel_K6": blatt["K6"].value,
        "diagramme": len(mappe[BLATT_PROJEKT]._charts),
    }


def _medien_reparieren(original: Path, ziel: Path) -> list[str]:
    """openpyxl verliert eingebettete Medien - fehlende Eintraege zurueckholen.

    Nachgewiesen fuer xl/media/image1.png (D-01). Der Vorgang ist bewusst
    allgemein gehalten: Er holt jeden Eintrag zurueck, den das Original hatte
    und die Kopie nicht mehr hat.
    """
    with zipfile.ZipFile(original) as a, zipfile.ZipFile(ziel) as b:
        fehlend = sorted(set(a.namelist()) - set(b.namelist()))
    if not fehlend:
        return []

    zwischen = ziel.with_suffix(".tmp")
    shutil.move(ziel, zwischen)
    with zipfile.ZipFile(original) as quelle, \
            zipfile.ZipFile(zwischen) as alt, \
            zipfile.ZipFile(ziel, "w", zipfile.ZIP_DEFLATED) as neu:
        for eintrag in alt.infolist():
            neu.writestr(eintrag, alt.read(eintrag.filename))
        for name in fehlend:
            neu.writestr(name, quelle.read(name))
    zwischen.unlink()
    return fehlend


def plan_erstellen(
    stammdaten: Stammdaten,
    berichte: list[Tagesbericht],
    blatt,
    *,
    bauvorhaben_erwartet: str | None = None,
    ueberschreiben: bool = False,
) -> Befuellungsplan:
    """Prueft alle Berichte und plant die Schreibvorgaenge - ohne zu schreiben."""
    plan = Befuellungsplan()

    if not stammdaten.bauzeit_gesetzt:
        plan.fehler.append(
            "Baubeginn und Bauende sind in der Mappe nicht gesetzt "
            "(Projektübersicht F11/F13). Ohne sie hat das Kalendergeruest "
            "keine Datumswerte und es gibt keine Zieladresse (M-1).")
        return plan

    if bauvorhaben_erwartet and stammdaten.bauvorhaben != bauvorhaben_erwartet:
        plan.fehler.append(
            f"Die Mappe gehoert zum Bauvorhaben {stammdaten.bauvorhaben!r}, "
            f"erwartet wurde {bauvorhaben_erwartet!r}. Abbruch, damit keine "
            "Stunden im falschen Projekt landen (D-04).")
        return plan

    bekannte_mitarbeiter = set(stammdaten.mitarbeiter)
    bekannte_gewerke = set(stammdaten.gewerke)

    # --- Berichte je Tag buendeln; Spalte I und J sind tagesweit ------------
    je_tag: dict[date, list[Tagesbericht]] = defaultdict(list)
    for bericht in berichte:
        je_tag[bericht.datum].append(bericht)

    for tag in sorted(je_tag):
        tagesberichte = sorted(je_tag[tag], key=lambda b: (b.mitarbeiter, b.bericht_id))

        # --- Tagesbezogene Pruefungen -------------------------------------
        if not ist_erfassungstag(tag, stammdaten.feiertage):
            grund = ("ein Sonntag - dafuer hat die Mappe keine Zeile"
                     if tag.weekday() == 6
                     else "ein Berliner Feiertag - die Eingabe ist gesperrt")
            plan.fehler.append(f"{tag:%d.%m.%Y} ist {grund} (D-06).")
            continue

        if stammdaten.bauende and tag > stammdaten.bauende:
            plan.fehler.append(
                f"{tag:%d.%m.%Y} liegt nach dem Bauende "
                f"{stammdaten.bauende:%d.%m.%Y}.")
            continue

        benoetigt = sum(b.benoetigte_slots for b in tagesberichte)
        if benoetigt > SLOTS_JE_TAG:
            namen = ", ".join(sorted({b.mitarbeiter for b in tagesberichte}))
            plan.fehler.append(
                f"{tag:%d.%m.%Y}: {benoetigt} Zeilen benoetigt, die Mappe "
                f"fasst {SLOTS_JE_TAG} je Tag. Betroffen: {namen} (R-03).")
            continue

        # --- Zeilen je Zeitfenster ----------------------------------------
        slot, tagesziel = 1, None
        tagesfehler = False
        for bericht in tagesberichte:
            if bericht.mitarbeiter not in bekannte_mitarbeiter:
                plan.fehler.append(
                    f"{tag:%d.%m.%Y}: {bericht.mitarbeiter!r} steht nicht in der "
                    "Mitarbeiterliste der Mappe. Die Stunden fielen sonst "
                    "lautlos aus der Auswertung (D-05).")
                tagesfehler = True
                continue

            for fenster in bericht.zeitfenster:
                if fenster.gewerk not in bekannte_gewerke:
                    plan.fehler.append(
                        f"{tag:%d.%m.%Y}: Gewerk {fenster.gewerk!r} steht nicht "
                        "in der Gewerkeliste der Mappe (D-05).")
                    tagesfehler = True
                    continue

                try:
                    ziel = datum_zu_zeile(stammdaten.baubeginn, tag, slot)
                except GeometrieFehler as fehler:
                    plan.fehler.append(str(fehler))
                    tagesfehler = True
                    continue

                tagesziel = tagesziel or ziel
                belegt = _belegte_zellen(blatt, ziel)
                if belegt and not ueberschreiben:
                    plan.fehler.append(
                        f"{tag:%d.%m.%Y}: Zeile {ziel.zeile} ist bereits "
                        f"befuellt ({', '.join(belegt)}). Mit --ueberschreiben "
                        "erzwingen (M-4).")
                    tagesfehler = True
                    continue

                plan.schreibvorgaenge += [
                    Schreibvorgang(ziel.mitarbeiter, bericht.mitarbeiter, "Mitarbeiter"),
                    Schreibvorgang(ziel.gewerk, fenster.gewerk, "Gewerk"),
                    Schreibvorgang(ziel.anfang, uhrzeit_zu_excel(fenster.beginn), "Arbeitsanfang"),
                    Schreibvorgang(ziel.ende, uhrzeit_zu_excel(fenster.ende), "Arbeitsende"),
                ]
                plan.zuordnungen.append(Zuordnung(
                    bericht_id=bericht.bericht_id, datum=tag,
                    mitarbeiter=bericht.mitarbeiter, zeile=ziel.zeile, slot=slot))
                slot += 1

        if tagesfehler or tagesziel is None:
            continue

        # --- Tagesweite Felder: Spalte I und J -----------------------------
        alle_leistungen = [l for b in tagesberichte for l in b.leistungen]
        if alle_leistungen:
            ids = [b.bericht_id for b in tagesberichte if b.leistungen]
            text = serialisieren(alle_leistungen, ids[0] if len(ids) == 1 else None)
            plan.schreibvorgaenge.append(
                Schreibvorgang(tagesziel.leistungen, text, "Leistungen des Tages"))
        else:
            plan.warnungen.append(
                f"{tag:%d.%m.%Y}: keine Leistungspositionen - Spalte I bleibt leer.")

        bemerkungen = [b.bemerkung.strip() for b in tagesberichte
                       if b.bemerkung and b.bemerkung.strip()]
        if bemerkungen:
            plan.schreibvorgaenge.append(Schreibvorgang(
                tagesziel.bemerkung, "; ".join(bemerkungen)[:500], "Tagesbemerkung"))

    return plan


def _belegte_zellen(blatt, ziel: Zielzellen) -> list[str]:
    return [zelle for zelle in (ziel.mitarbeiter, ziel.gewerk, ziel.anfang, ziel.ende)
            if blatt[zelle].value not in (None, "")]


def mappe_befuellen(
    quelle: Path,
    ziel: Path,
    berichte: list[Tagesbericht],
    *,
    bauvorhaben_erwartet: str | None = None,
    ueberschreiben: bool = False,
    probelauf: bool = False,
) -> Befuellungsplan:
    """Befuellt eine Kopie der Mappe. Die Quelle wird nur gelesen (D-01)."""
    # Die Mappe wird genau einmal geladen: Stammdaten, Geometriepruefung,
    # Planung und Schreiben arbeiten alle auf derselben Instanz.
    mappe = openpyxl.load_workbook(quelle)
    stammdaten = stammdaten_aus_mappe(mappe)      # prueft auch die Geometrie
    kennzahlen_vorher = _kennzahlen(mappe)
    blatt = mappe[BLATT_ZEITERFASSUNG]

    plan = plan_erstellen(
        stammdaten, berichte, blatt,
        bauvorhaben_erwartet=bauvorhaben_erwartet, ueberschreiben=ueberschreiben)

    if probelauf or not plan.ist_ausfuehrbar:
        mappe.close()
        return plan

    for vorgang in plan.schreibvorgaenge:
        blatt[vorgang.zelle] = vorgang.wert

    ziel.parent.mkdir(parents=True, exist_ok=True)
    mappe.save(ziel)
    mappe.close()

    repariert = _medien_reparieren(quelle, ziel)
    if repariert:
        plan.warnungen.append(
            f"Medien nach dem Schreiben wiederhergestellt: {', '.join(repariert)}")

    # --- Nachpruefung: ist die Mappe noch intakt? --------------------------
    geschrieben = openpyxl.load_workbook(ziel)
    kennzahlen_nachher = _kennzahlen(geschrieben)
    geschrieben.close()
    abweichungen = {k: (kennzahlen_vorher[k], kennzahlen_nachher[k])
                    for k in kennzahlen_vorher
                    if kennzahlen_vorher[k] != kennzahlen_nachher[k]}
    if abweichungen:
        ziel.unlink(missing_ok=True)
        plan.fehler.append(
            f"Die Mappe hat das Schreiben nicht unbeschadet ueberstanden: "
            f"{abweichungen}. Ausgabe verworfen.")
    return plan
